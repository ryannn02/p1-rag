"""解析与切分：把 data/raw 里的文件转成 data/parsed/*.txt 和 data/chunks.jsonl。

    python -m scripts.build_corpus

切分策略：优先按「章 / 节 / 条」切，条款不跨块；无条款编号的文档按段落聚合到上限长度。
"""

import argparse
import csv
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from bs4 import BeautifulSoup

from scripts.probe import FIELDS

RAW = Path("data/raw")
PARSED = Path("data/parsed")
CHUNKS = Path("data/chunks.jsonl")
META = Path("meta/documents.csv")
FAILURES = Path("reports/parse_failures.csv")

MAX_CHARS = 800
MIN_CHUNK = 50
MIN_CHARS = 20
MIN_DOC_CHARS = 120

HEADING = re.compile(r"^\s*第\s*[一二三四五六七八九十百零〇\d]+\s*[条章节编节]|^\s*第\s*[一二三四五六七八九十]+\s*部分")
NOISE = re.compile(r"网站总访问人次|版权所有|沪ICP备|浏览次数|加入收藏|设为首页|友情链接|"
                   r"^\s*发布时间[:：]|^\s*来源[:：]|^\s*打印本页|^\s*关闭窗口|^\s*分享到")
PAGENO = re.compile(r"^\s*[-—\s]*\d{1,3}\s*[-—\s]*$")
HTML_SELECTORS = [".wp_articlecontent", "div.mm", ".TRS_Editor", ".v_news_content",
                  "#content", ".article-content", ".Article_Content", ".entry"]

csv.field_size_limit(10**7)


def join_short_lines(text, thresh=25):
    """HTML 里 <span> 拆出来的碎行合并回一句，否则「根据 / 《财政部 / 教育部」会各占一行。"""
    out = []
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        mergeable = (
            out
            and len(line) < thresh
            and not HEADING.match(line)
            and not out[-1].endswith(("。", "！", "？", "；", "："))
        )
        if mergeable:
            out[-1] += line
        else:
            out.append(line)
    return "\n".join(out)


def clean_text(text):
    lines = []
    for raw in (text or "").replace("\r", "\n").split("\n"):
        line = re.sub(r"[ \t\u3000]+", " ", raw).strip()
        if not line or NOISE.search(line) or PAGENO.match(line):
            continue
        lines.append(line)
    out, prev = [], None
    for line in lines:
        if line != prev:
            out.append(line)
        prev = line
    return join_short_lines("\n".join(out))


def parse_pdf(path):
    import fitz

    doc = fitz.open(str(path))
    parts = []
    for i, page in enumerate(doc):
        parts.append((i + 1, page.get_text() or ""))
    doc.close()
    if not parts:
        return []
    text = "\n".join(t for _, t in parts)
    if len(text.strip()) < MIN_DOC_CHARS:
        return []
    pages = []
    cursor = 0
    for pno, t in parts:
        start = cursor
        cursor += len(t) + 1
        pages.append((pno, start, cursor))
    return [(text, pages)]


def parse_docx(path):
    import docx

    d = docx.Document(str(path))
    out = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    text = "\n".join(out)
    # python-docx 读不到文本框、图形里的文字，这些制度常常整篇塞在文本框里，
    # 所以再直接扫一遍 document.xml 兜底。
    if len(text.strip()) < 200:
        raw = parse_docx_xml(path)
        if len(raw.strip()) > len(text.strip()):
            return [(raw, None)]
    return [(text, None)]


def parse_docx_xml(path):
    import html as htmllib
    import zipfile

    with zipfile.ZipFile(str(path)) as z:
        names = [n for n in z.namelist() if n.startswith("word/") and n.endswith(".xml")]
        chunks = [z.read(n).decode("utf-8", "ignore") for n in names]
    xml = "\n".join(chunks)
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<w:tab[^>]*/>", "\t", xml)
    return htmllib.unescape(re.sub(r"<[^>]+>", "", xml))


def _doc_text(path):
    """`.doc` 是二进制格式，没有纯 Python 解析库，借系统工具转文本。

    macOS 自带 `textutil`（能直接写 stdout）；Windows / Linux 装 LibreOffice 后有 `soffice`，
    但它只能写文件——转进临时目录再读回，并显式指定 UTF-8，免得落到系统默认编码
    （Windows 上是 GBK，中文会全烂）。
    """
    if shutil.which("textutil"):
        r = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(path)],
                           capture_output=True, timeout=60)
        return r.stdout.decode("utf-8", errors="ignore")
    if shutil.which("soffice"):
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(["soffice", "--headless", "--convert-to",
                            "txt:Text (encoded):UTF8", "--outdir", tmp, str(path)],
                           capture_output=True, timeout=180)
            out = Path(tmp) / f"{path.stem}.txt"
            return out.read_text(encoding="utf-8", errors="ignore") if out.exists() else ""
    print("    未找到 .doc 转换工具，跳过该文件：macOS 用自带的 textutil；"
          "Windows / Linux 装 LibreOffice 即可", file=sys.stderr)
    return ""


def parse_doc(path):
    try:
        text = _doc_text(path)
    except Exception as exc:
        print(f"    .doc 转换失败：{exc}", file=sys.stderr)
        return []
    return [(text, None)] if text.strip() else []


def parse_txt(path):
    return [(path.read_text(encoding="utf-8", errors="ignore"), None)]


def parse_html(path):
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    best = ""
    for sel in HTML_SELECTORS:
        node = soup.select_one(sel)
        if node:
            t = node.get_text("\n", strip=True)
            if len(t) > len(best):
                best = t
    if len(best) < MIN_DOC_CHARS:
        best = max((d.get_text("\n", strip=True) for d in soup.find_all("div")), key=len, default="")
    return [(best, None)]


PARSERS = {"pdf": parse_pdf, "docx": parse_docx, "doc": parse_doc, "html": parse_html,
           "htm": parse_html, "txt": parse_txt}


def split_chunks(text, pages):
    """按条款切分；没有条款编号的按段落聚合。返回 [(section, page, text)]"""
    lines = [l for l in text.split("\n")]
    blocks, section = [], []
    stack = []
    for line in lines:
        if HEADING.match(line) and len(line) < 60:
            head = re.match(r"^\s*(第\s*[一二三四五六七八九十百零〇\d]+\s*[条章节编节])", line)
            if section:
                blocks.append((" / ".join(stack), "\n".join(section)))
            if head and "章" in head.group(1):
                stack = [line.strip()]
            elif head and "节" in head.group(1) and len(stack) <= 1:
                stack = stack[:1] + [line.strip()]
            else:
                stack = stack or []
            section = [line] if not head or "条" in head.group(1) else [line]
        else:
            section.append(line)
    if section:
        blocks.append((" / ".join(stack), "\n".join(section)))

    out = []
    for sec, body in blocks:
        body = body.strip()
        if len(body) < MIN_CHARS:
            continue
        page = None
        if pages:
            pos = text.find(body[:40])
            for pno, start, end in pages:
                if start <= pos < end:
                    page = pno
                    break
        for piece in hard_split(body):
            out.append((sec, page, piece))
    return out


def hard_split(body):
    """超长块切分：先按句号拆开长段落，再贪心聚合回上限长度。"""
    pieces = []
    for para in body.split("\n"):
        if len(para) <= MAX_CHARS:
            pieces.append(para)
            continue
        sentences = re.split(r"(?<=[。；！？])", para)
        buf = ""
        for sent in sentences:
            if len(buf) + len(sent) > MAX_CHARS and buf:
                pieces.append(buf)
                buf = sent
            else:
                buf += sent
        if buf:
            pieces.append(buf)

    out, cur = [], ""
    for piece in pieces:
        if cur and len(cur) + len(piece) + 1 > MAX_CHARS:
            out.append(cur)
            cur = piece
        else:
            cur = f"{cur}\n{piece}" if cur else piece
    if cur:
        out.append(cur)

    final = []
    for piece in out:
        piece = piece.strip()
        while len(piece) > MAX_CHARS:
            final.append(piece[:MAX_CHARS])
            piece = piece[MAX_CHARS:]
        if piece:
            final.append(piece)
    return final


def merge_small(chunks):
    """碎片（页眉页脚、表格残渣）并进上一块；并不动又太短的丢掉。"""
    out = []
    for c in chunks:
        if out and c["chars"] < MIN_CHUNK and out[-1]["chars"] + c["chars"] <= MAX_CHARS:
            out[-1]["text"] += "\n" + c["text"]
            out[-1]["chars"] = len(out[-1]["text"])
        elif c["chars"] >= MIN_CHUNK:
            out.append(c)
    return out


def prune_empty(failures):
    """网页正文区为空的条目在官网上就是空壳，永远解析不出内容，直接标 excluded。

    只处理 html：图片型 PDF 现在没文字层，但 OCR 后还有救，不能一并丢掉。
    """
    empty = [f for f in failures
             if f["reason"].startswith("text_too_short") and f["format"] in ("html", "htm")]
    if not empty:
        print("没有正文为空的网页条目")
        return
    ids = {f["doc_id"] for f in empty}
    rows = list(csv.DictReader(META.open(encoding="utf-8")))
    for r in rows:
        if r["doc_id"] in ids:
            r["status"] = "excluded"
    with META.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)
    print(f"已标记 {len(ids)} 条正文为空的网页为 excluded：")
    for f in empty:
        print(f"  {f['dept']:<6}{f['title'][:44]}")


def clear_parsed():
    """重建前清掉旧产物：分诊剔掉的文档不该在 data/parsed 里留残骸。"""
    n = 0
    for f in PARSED.glob("*.txt"):
        f.unlink()
        n += 1
    if n:
        print(f"清掉上一轮 {n} 个全文 txt")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--prune-empty", action="store_true",
                    help="把正文为空的网页条目标成 excluded（写入 meta/documents.csv）")
    args = ap.parse_args()

    PARSED.mkdir(parents=True, exist_ok=True)
    FAILURES.parent.mkdir(parents=True, exist_ok=True)
    clear_parsed()
    # archived 是已被新版替代的历史版本，保留在册供追溯，但不进检索
    rows = [r for r in csv.DictReader(META.open(encoding="utf-8")) if r["status"] == "effective"]
    print(f"待解析 {len(rows)} 条")

    failures, chunk_rows, n_chunks = [], [], 0
    for r in rows:
        path = Path(r["local_path"])
        parser = PARSERS.get(r["format"])
        if not path.exists() or parser is None:
            failures.append({**r, "reason": "file_missing_or_unsupported_format"})
            continue
        try:
            parts = parser(path)
        except Exception as exc:
            failures.append({**r, "reason": f"parse_error: {exc}"})
            continue

        text, pages = (parts[0] if parts else ("", None))
        text = clean_text(text)
        if len(text) < MIN_DOC_CHARS:
            failures.append({**r, "reason": f"text_too_short({len(text)})"})
            continue

        (PARSED / f"{r['doc_id']}.txt").write_text(text, encoding="utf-8")
        doc_chunks = []
        for sec, page, piece in split_chunks(text, pages):
            doc_chunks.append({
                "chunk_id": f"{r['doc_id']}-{n_chunks:05d}",
                "doc_id": r["doc_id"], "dept": r["dept"], "title": r["title"],
                "year": r["year"], "section": sec, "page": page or "",
                "text": piece, "chars": len(piece), "source_url": r["source_url"],
            })
        for c in merge_small(doc_chunks):
            c["chunk_id"] = f"{r['doc_id']}-{n_chunks:05d}"
            chunk_rows.append(c)
            n_chunks += 1

    with CHUNKS.open("w", encoding="utf-8") as f:
        for c in chunk_rows:
            f.write(json.dumps(c, ensure_ascii=False) + "\n")

    if failures:
        with FAILURES.open("w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(failures[0].keys()))
            w.writeheader()
            w.writerows(failures)

    if not chunk_rows and failures:
        print("\n没有解析出任何内容。若 data/raw 是空的，请先从共享位置同步原始文件再跑；"
              "若确实有文件，检查 reports/parse_failures.csv 的原因列。", file=sys.stderr)
        return 2

    if args.prune_empty:
        prune_empty(failures)

    ok = len(rows) - len(failures)
    lens = [c["chars"] for c in chunk_rows] or [0]
    print(f"解析成功 {ok} 条，失败 {len(failures)} 条（明细见 {FAILURES}）")
    print(f"切出 {n_chunks} 个 chunk，平均 {sum(lens)//len(lens)} 字，最长 {max(lens)} 字")
    print("失败原因分布：")
    for reason, n in sorted({f["reason"].split("(")[0]: 0 for f in failures}.items()):
        cnt = sum(1 for f in failures if f["reason"].split("(")[0] == reason)
        print(f"  {reason:<34}{cnt}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
